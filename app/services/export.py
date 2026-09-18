from __future__ import annotations

from io import BytesIO
import base64

from docx import Document
from docx.shared import Inches

from app.models.domain import Article
from app.services.illustrations import inline_content


def markdown(article: Article) -> str:
    parts = [f"# {article.title}"]
    for paragraph, images in inline_content(article):
        parts.append(paragraph)
        for picture in images:
            parts.append(f"![{picture.label}](data:image/jpeg;base64,{picture.image_b64})")
    return "\n\n".join(parts) + f"\n\n---\n生成任务：{article.task_id}\n"


def plain_text(article: Article) -> str:
    return f"{article.title}\n\n{article.body}"


def word_document(article: Article) -> bytes:
    document = Document()
    document.add_heading(article.title, 0)
    for paragraph, images in inline_content(article):
        document.add_paragraph(paragraph)
        for picture in images:
            document.add_picture(BytesIO(base64.b64decode(picture.image_b64)), width=Inches(5.5))
            document.add_paragraph(picture.label)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
